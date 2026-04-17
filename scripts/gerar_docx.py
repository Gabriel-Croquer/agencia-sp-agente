"""
Converte materia em Markdown para Word (.docx) com formatacao da Agencia SP.
Uso: python scripts/gerar_docx.py <arquivo.md>
Gera o .docx no mesmo diretorio do .md original.
"""

import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def criar_estilos(doc):
    """Configura estilos base do documento."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(3)


def adicionar_titulo(doc, texto):
    """Adiciona titulo H1."""
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    p.paragraph_format.space_after = Pt(4)


def adicionar_linha_fina(doc, texto):
    """Adiciona linha fina / meta description."""
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.italic = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p.paragraph_format.space_after = Pt(16)


def adicionar_tags(doc, texto):
    """Adiciona tags em fonte menor."""
    p = doc.add_paragraph()
    run = p.add_run(f"Tags: {texto}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    p.paragraph_format.space_after = Pt(20)


def adicionar_intertitulo(doc, texto):
    """Adiciona intertitulo H2."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)


def adicionar_paragrafo(doc, texto):
    """Adiciona paragrafo normal."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(texto)
    run.font.size = Pt(11)


def adicionar_fonte(doc, texto):
    """Adiciona credito de fonte em italico."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run(texto)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def adicionar_infografico(doc, texto):
    """Adiciona bloco de sugestao de infografico."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    run_label = p.add_run("SUGESTAO DE INFOGRAFICO: ")
    run_label.bold = True
    run_label.font.size = Pt(10)
    run_label.font.color.rgb = RGBColor(0x00, 0x66, 0x99)
    run_text = p.add_run(texto)
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def parse_md(caminho_md):
    """Le o Markdown e converte para Word."""
    texto = Path(caminho_md).read_text(encoding="utf-8")
    linhas = texto.split("\n")

    doc = Document()
    criar_estilos(doc)

    meta_desc = ""
    tags = ""
    bloco_infografico = []
    em_infografico = False
    bloco_fonte = []
    em_fonte = False

    i = 0
    while i < len(linhas):
        linha = linhas[i].strip()

        # Pular separadores
        if linha == "---":
            i += 1
            continue

        # Titulo H1
        if linha.startswith("# ") and not linha.startswith("## "):
            adicionar_titulo(doc, linha[2:].strip())
            i += 1
            continue

        # Meta description
        if linha.startswith("**Meta description:**"):
            meta_desc = linha.replace("**Meta description:**", "").strip()
            adicionar_linha_fina(doc, meta_desc)
            i += 1
            continue

        # Tags
        if linha.startswith("**Tags:**"):
            tags = linha.replace("**Tags:**", "").strip()
            adicionar_tags(doc, tags)
            i += 1
            continue

        # Intertitulo H2
        if linha.startswith("## "):
            adicionar_intertitulo(doc, linha[3:].strip())
            i += 1
            continue

        # Sugestao de infografico
        if linha.startswith("**Sugestão de infográfico:**") or linha.startswith("**Sugestao de infografico:**"):
            texto_info = re.sub(r"\*\*Sugest[ãa]o de infogr[áa]fico:\*\*\s*", "", linha)
            # Coletar linhas seguintes se o bloco continuar
            j = i + 1
            while j < len(linhas) and linhas[j].strip() and not linhas[j].strip().startswith("---"):
                texto_info += " " + linhas[j].strip()
                j += 1
            adicionar_infografico(doc, texto_info)
            i = j
            continue

        # Fonte (italico no final)
        if linha.startswith("*Fonte:") or linha.startswith("*fonte:"):
            texto_fonte = linha.strip("*").strip()
            adicionar_fonte(doc, texto_fonte)
            i += 1
            continue

        # Paragrafo normal
        if linha:
            adicionar_paragrafo(doc, linha)

        i += 1

    # Salvar
    caminho_docx = Path(caminho_md).with_suffix(".docx")
    doc.save(str(caminho_docx))
    return caminho_docx


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python scripts/gerar_docx.py <arquivo.md>")
        sys.exit(1)

    caminho = sys.argv[1]
    resultado = parse_md(caminho)
    print(f"Arquivo gerado: {resultado}")
